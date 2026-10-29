from django.contrib.auth.decorators import login_required
from django.conf import settings
from django.db.models import Count, Q
import csv
from datetime import date, datetime
from pathlib import Path
from django.contrib import messages
from django.http import HttpResponse, HttpResponseForbidden
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.shortcuts import get_object_or_404, redirect, render
from django.forms import inlineformset_factory
from accounts.access import is_system_administrator, record_editor_required, system_admin_required
from .forms import FieldAttendanceForm, FieldAttendanceParticipantForm
from .models import FieldAttendance, FieldAttendanceParticipant, initials_from_name

from .models import School
from projects.models import Indicator
from monitoring.models import FieldVisit, MHMSession
from school_governance.models import GovernanceMeeting, GovernanceTraining
from learners.models import AttendanceRecord


ATTENDANCE_THIN_GRAY = Side(style="thin", color="D9D9D9")


@login_required
def school_detail(request, pk):
    school = get_object_or_404(School.objects.select_related("location"), pk=pk)
    enrolments = school.enrolments.select_related("learner", "project", "academic_year").order_by("-academic_year__start_date", "learner__last_name")
    current_enrolments = enrolments.filter(is_current=True)
    return render(request, "core/school_detail.html", {
        "school": school,
        "enrolments": enrolments[:25],
        "enrolled_count": current_enrolments.count(),
        "total_learner_count": school.enrolments.values("learner_id").distinct().count(),
        "status_counts": current_enrolments.values("learner__status").annotate(total=Count("learner_id")).order_by("learner__status"),
        "gender_counts": current_enrolments.values("learner__sex").annotate(total=Count("learner_id")).order_by("learner__sex"),
        "participant_counts": current_enrolments.values("learner__participant_category").annotate(total=Count("learner_id")).order_by("learner__participant_category"),
        "displacement_counts": current_enrolments.values("learner__displacement_status").annotate(total=Count("learner_id")).order_by("learner__displacement_status"),
        "project_counts": enrolments.values("project__name").annotate(total=Count("learner_id", distinct=True)).order_by("project__name"),
        "activities": Indicator.objects.filter(project__isnull=False, data_entries__school=school).select_related("project").distinct().order_by("-created_at")[:8],
        "school_attendance_count": AttendanceRecord.objects.filter(enrolment__school=school).count(),
        "field_visits": FieldVisit.objects.filter(school=school).select_related("project").order_by("-visit_date")[:8],
        "mhm_sessions": MHMSession.objects.filter(school=school).select_related("project").order_by("-session_date")[:8],
        "wash_facilities": school.wash_facilities.select_related("project").order_by("facility_type"),
        "governance_members": school.governance_members.filter(active=True).order_by("body", "full_name"),
        "governance_meetings": GovernanceMeeting.objects.filter(school=school).order_by("-meeting_date")[:8],
        "governance_trainings": GovernanceTraining.objects.filter(school=school).order_by("-training_date")[:8],
    })


ParticipantFormSet = inlineformset_factory(
    FieldAttendance, FieldAttendanceParticipant,
    form=FieldAttendanceParticipantForm,
    extra=3, max_num=25, can_delete=True,
)


@login_required
def attendance_list(request):
    sheets, participant_filter = _attendance_filtered_queryset(request)
    sheets = sheets.annotate(
        participant_total=Count("participants", filter=participant_filter, distinct=True),
        male_total=Count("participants", filter=participant_filter & Q(participants__sex="M"), distinct=True),
        female_total=Count("participants", filter=participant_filter & Q(participants__sex="F"), distinct=True),
    )
    totals = sheets.aggregate(
        male=Count("participants", filter=participant_filter & Q(participants__sex="M"), distinct=True),
        female=Count("participants", filter=participant_filter & Q(participants__sex="F"), distinct=True),
        total=Count("participants", filter=participant_filter, distinct=True),
    )
    return render(request, "core/attendance_list.html", {
        "sheets": sheets,
        "can_review": is_system_administrator(request.user),
        "activities": Indicator.objects.order_by("name"),
        "status_choices": FieldAttendance._meta.get_field("status").choices,
        "sex_choices": FieldAttendanceParticipant._meta.get_field("sex").choices,
        "filters": request.GET,
        "filter_query": request.GET.urlencode(),
        "totals": totals,
    })


@login_required
def attendance_detail(request, pk):
    sheet = get_object_or_404(
        FieldAttendance.objects.select_related("submitted_by", "reviewed_by").prefetch_related("participants"),
        pk=pk,
    )
    return render(request, "core/attendance_detail.html", {
        "sheet": sheet,
        "can_review": is_system_administrator(request.user),
    })


def _attendance_form(request, instance=None):
    form = FieldAttendanceForm(request.POST or None, instance=instance)
    formset = ParticipantFormSet(request.POST or None, instance=instance, prefix="participants")
    if request.method == "POST" and form.is_valid() and formset.is_valid():
        sheet = form.save(commit=False)
        reopening_approved = bool(instance and instance.status == "approved" and is_system_administrator(request.user))
        if not sheet.submitted_by_id:
            sheet.submitted_by = request.user
        if not sheet.submitted_date:
            sheet.submitted_date = timezone.localdate()
        if not sheet.submitted_signature:
            sheet.submitted_signature = initials_from_name(request.user.get_full_name() or request.user.username)
            sheet.submitted_title = request.user.job_title
        if reopening_approved:
            sheet.status = "pending"
            sheet.reviewed_by = None
            sheet.reviewed_at = None
            sheet.approved_by = ""
            sheet.approved_signature = ""
            sheet.approved_date = None
            sheet.rejection_reason = ""
        sheet.save()
        formset.instance = sheet
        formset.save()
        messages.success(request, "Field attendance sheet saved successfully.")
        return redirect("data:field_attendance_detail", pk=sheet.pk)
    return render(request, "core/attendance_form.html", {"form": form, "formset": formset, "object": instance})


@login_required
@record_editor_required
def attendance_create(request):
    return _attendance_form(request)


@login_required
@record_editor_required
def attendance_update(request, pk):
    instance = get_object_or_404(FieldAttendance, pk=pk)
    if instance.status == "approved" and not is_system_administrator(request.user):
        return HttpResponseForbidden("Approved attendance records cannot be edited.")
    return _attendance_form(request, instance)


@login_required
@system_admin_required
def attendance_approve(request, pk):
    if request.method == "POST":
        sheet = get_object_or_404(FieldAttendance, pk=pk)
        sheet.status = "approved"
        sheet.reviewed_by = request.user
        sheet.reviewed_at = timezone.now()
        sheet.approved_by = request.user.get_full_name() or request.user.username
        sheet.approved_signature = initials_from_name(sheet.approved_by)
        sheet.approved_date = timezone.localdate()
        sheet.rejection_reason = ""
        sheet.save(update_fields=["status", "reviewed_by", "reviewed_at", "approved_by", "approved_signature", "approved_date", "rejection_reason", "updated_at"])
        messages.success(request, "Field attendance approved.")
    return redirect("data:field_attendance_detail", pk=pk)


@login_required
@system_admin_required
def attendance_reject(request, pk):
    if request.method == "POST":
        sheet = get_object_or_404(FieldAttendance, pk=pk)
        sheet.status = "rejected"
        sheet.reviewed_by = request.user
        sheet.reviewed_at = timezone.now()
        sheet.rejection_reason = request.POST.get("rejection_reason", "").strip()
        sheet.save(update_fields=["status", "reviewed_by", "reviewed_at", "rejection_reason", "updated_at"])
        messages.warning(request, "Field attendance rejected and returned for correction.")
    return redirect("data:field_attendance_detail", pk=pk)


def _attendance_filtered_queryset(request, approved_only=False):
    sheets = FieldAttendance.objects.all()
    if approved_only:
        sheets = sheets.filter(status="approved")

    central_filter_request = any(key.startswith("attendance_") for key in request.GET)

    def attendance_param(name):
        if f"attendance_{name}" in request.GET:
            return request.GET.get(f"attendance_{name}", "").strip()
        if central_filter_request:
            return ""
        return request.GET.get(name, "").strip()

    activity = attendance_param("activity")
    date_from = attendance_param("date_from")
    date_to = attendance_param("date_to")
    district = attendance_param("district")
    status = attendance_param("status")
    sex = attendance_param("sex")
    participant = attendance_param("participant")
    if activity:
        sheets = sheets.filter(activity_id=activity)
    if date_from:
        sheets = sheets.filter(activity_date__gte=date_from)
    if date_to:
        sheets = sheets.filter(activity_date__lte=date_to)
    if district:
        sheets = sheets.filter(district__icontains=district)
    if status:
        sheets = sheets.filter(status=status)

    participant_filter = Q()
    if sex:
        participant_filter &= Q(participants__sex=sex)
    if participant:
        participant_filter &= Q(participants__name__icontains=participant)
    if sex or participant:
        sheets = sheets.filter(participant_filter)
    return sheets.prefetch_related("participants").select_related("submitted_by", "reviewed_by", "activity"), participant_filter


def _attendance_export_rows(request):
    sheets, _ = _attendance_filtered_queryset(request)
    central_filter_request = any(key.startswith("attendance_") for key in request.GET)
    sex = request.GET.get("attendance_sex", "").strip()
    participant_name = request.GET.get("attendance_participant", "").strip().lower()
    if not central_filter_request:
        sex = request.GET.get("sex", "").strip()
        participant_name = request.GET.get("participant", "").strip().lower()
    totals = {"male": 0, "female": 0, "total": 0}
    for sheet in sheets:
        for participant in sheet.participants.all():
            if (sex and participant.sex != sex) or (participant_name and participant_name not in participant.name.lower()):
                continue
            totals["total"] += 1
            if participant.sex == "M":
                totals["male"] += 1
            elif participant.sex == "F":
                totals["female"] += 1
            submitted_name = sheet.submitted_by.get_full_name() if sheet.submitted_by_id else ""
            yield [sheet.activity_date, str(sheet.activity or ""), sheet.district, sheet.county, sheet.sub_county, sheet.parish, sheet.village, submitted_name, sheet.submitted_signature, sheet.approved_by, sheet.approved_signature, participant.name, participant.sex, participant.title_role, participant.contact, participant.signature]
    yield []
    yield ["Summary", "No. of male", totals["male"]]
    yield ["Summary", "No. of female", totals["female"]]
    yield ["Summary", "Total participants", totals["total"]]


def _attendance_filter_description(request):
    labels = []
    central_filter_request = any(key.startswith("attendance_") for key in request.GET)
    for key, label in (("activity", "Activity"), ("date_from", "From"), ("date_to", "To"), ("district", "District"), ("status", "Status"), ("sex", "Sex"), ("participant", "Participant")):
        if f"attendance_{key}" in request.GET:
            value = request.GET.get(f"attendance_{key}", "").strip()
        elif central_filter_request:
            value = ""
        else:
            value = request.GET.get(key, "").strip()
        if value:
            labels.append(f"{label}: {value}")
    return "; ".join(labels) or "All attendance records"


def _add_attendance_report_header(sheet, request, headers):
    navy = "17365D"
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    sheet.merge_cells(start_row=3, start_column=1, end_row=3, end_column=len(headers))
    sheet["A1"] = "KAPDA ATTENDANCE SHEET"
    sheet["A2"] = "Field attendance report"
    sheet["A3"] = f"Filters: {_attendance_filter_description(request)}"
    sheet["A1"].font = Font(name="Arial", size=16, bold=True, color=navy)
    sheet["A2"].font = Font(name="Arial", size=11, bold=True, color=navy)
    sheet["A3"].font = Font(name="Arial", size=9, italic=True, color="666666")
    for cell_ref in ("A1", "A2", "A3"):
        sheet[cell_ref].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[2].height = 22
    sheet.row_dimensions[3].height = 20
    sheet.row_dimensions[5].height = 28
    sheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=len(headers))
    sheet["A5"] = "Attendance records"
    sheet["A5"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    sheet["A5"].fill = PatternFill("solid", fgColor=navy)
    sheet["A5"].alignment = Alignment(horizontal="left", vertical="center")
    for index, header in enumerate(headers, 1):
        cell = sheet.cell(row=6, column=index, value=header)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=ATTENDANCE_THIN_GRAY, bottom=ATTENDANCE_THIN_GRAY, left=ATTENDANCE_THIN_GRAY, right=ATTENDANCE_THIN_GRAY)
    sheet.freeze_panes = "A7"
    sheet.sheet_view.showGridLines = False
    logo_dir = Path(settings.BASE_DIR) / "static" / "images"
    for filename, anchor, width in (("logo.png", "A1", 85), ("mcc-logo.jpg", f"{get_column_letter(len(headers))}1", 105)):
        image_path = logo_dir / filename
        if image_path.exists():
            image = ExcelImage(str(image_path))
            image.width = width
            image.height = 45
            sheet.add_image(image, anchor)


@login_required
def attendance_export_csv(request):
    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="field_attendance.csv"'
    writer = csv.writer(response)
    writer.writerow(["KAPDA ATTENDANCE SHEET"])
    writer.writerow(["Field attendance report"])
    writer.writerow(["KAPDA logo", "static/images/logo.png", "MCC logo", "static/images/mcc-logo.jpg"])
    writer.writerow([f"Filters: {_attendance_filter_description(request)}"])
    writer.writerow([])
    writer.writerow(["Activity date", "Activity", "District", "County", "Sub-county", "Parish", "Village", "Submitted by", "Submitted initials", "Approved by", "Approved initials", "Participant name", "Sex", "Title / role", "Email / telephone", "Participant initials"])
    writer.writerows(_attendance_export_rows(request))
    return response


@login_required
def attendance_export_xlsx(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Field attendance"
    headers = ["Activity date", "Activity", "District", "County", "Sub-county", "Parish", "Village", "Submitted by", "Submitted initials", "Approved by", "Approved initials", "Participant name", "Sex", "Title / role", "Email / telephone", "Participant initials"]
    _add_attendance_report_header(sheet, request, headers)
    light_blue = "D9EAF7"
    for row in _attendance_export_rows(request):
        sheet.append(row)
    for row in sheet.iter_rows(min_row=7):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(top=ATTENDANCE_THIN_GRAY, bottom=ATTENDANCE_THIN_GRAY, left=ATTENDANCE_THIN_GRAY, right=ATTENDANCE_THIN_GRAY)
            if cell.column == 1 and isinstance(cell.value, (date, datetime)):
                cell.number_format = "dd/mm/yyyy"
        if row[0].value == "Summary":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=light_blue)
                cell.font = Font(name="Arial", size=10, bold=True, color="17365D")
    sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{sheet.max_row - 4}"
    for column in sheet.columns:
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(max(len(str(cell.value or "")) for cell in column) + 2, 32)
    sheet.column_dimensions["A"].width = 15
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="field_attendance.xlsx"'
    workbook.save(response)
    return response
