from datetime import date
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Count, Max, Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse
from django.utils.dateparse import parse_date
from core.models import AcademicYear, AuditEvent, School
from .models import Learner, LearnerProject, Enrolment, AttendanceRecord, ClassMovement, LearnerSupport, UNEBResult, LearnerEmployment, LearnerStatusHistory
from monitoring.models import FollowUpAction, MHMSession, WASHFacility
from .forms import LearnerForm, EnrolmentForm, LearnerSupportForm, LearnerEmploymentForm, LearnerStatusHistoryForm
from accounts.access import record_editor_required
from projects.models import Project

@login_required
def learner_list(request):
    qs = Learner.objects.select_related("location").all()
    q = request.GET.get("q","").strip()
    project = request.GET.get("project","")
    class_level = request.GET.get("class","")
    status = request.GET.get("status","")
    participant_category = request.GET.get("participant_category","")
    displacement_status = request.GET.get("displacement_status","")
    support_type = request.GET.get("support_type","")
    if q: qs=qs.filter(Q(reg_no__icontains=q)|Q(first_name__icontains=q)|Q(middle_name__icontains=q)|Q(last_name__icontains=q)|Q(guardian_phone__icontains=q))
    if class_level: qs=qs.filter(current_class=class_level)
    if status: qs=qs.filter(status=status)
    if participant_category: qs=qs.filter(participant_category=participant_category)
    if displacement_status: qs=qs.filter(displacement_status=displacement_status)
    if support_type: qs=qs.filter(supports__support_type=support_type)
    if project: qs=qs.filter(project_links__project_id=project,project_links__active=True)
    qs=qs.distinct()
    learner_ids = qs.values("pk")
    support_scope = LearnerSupport.objects.filter(learner_id__in=learner_ids)
    if project:
        support_scope = support_scope.filter(project_id=project)
    supported_learners = Learner.objects.filter(pk__in=support_scope.values("learner_id")).distinct()
    girls_women_reached = supported_learners.filter(Q(sex="F") | Q(participant_category__in=["girl", "woman"])).count()
    school_support = list(
        support_scope.filter(learner__enrolments__is_current=True)
        .values("learner__enrolments__school_id", "learner__enrolments__school__name")
        .annotate(total=Count("learner_id", distinct=True))
        .order_by("learner__enrolments__school__name")
    )
    for row in school_support:
        row["school_url"] = reverse(
            "learners:supported_learners_by_school",
            args=[row["learner__enrolments__school_id"]],
        )
    school_project_support = school_support
    activity_project_filter = {"project_id": project} if project else {}
    follow_up_scope = FollowUpAction.objects.all()
    if project:
        follow_up_scope = follow_up_scope.filter(Q(visit__project_id=project) | Q(activity__project_id=project)).distinct()
    kpis = {
        "total_learners": qs.count(),
        "supported_learners": supported_learners.count(),
        "girls_women_reached": girls_women_reached,
        "pad_distribution": support_scope.filter(support_type="sanitary_pads_mhm").count(),
        "back_to_school": support_scope.filter(support_type="back_to_school").count(),
        "mhm_sessions": MHMSession.objects.filter(**activity_project_filter).count(),
        "wash_facilities": WASHFacility.objects.filter(**activity_project_filter).count(),
        "follow_up": follow_up_scope.count(),
    }
    page=Paginator(qs,25).get_page(request.GET.get("page"))
    return render(request,"learners/learner_dashboard.html",{"learners":page,"q":q,"project_filter":project,"class_filter":class_level,"status_filter":status,"participant_category_filter":participant_category,"displacement_status_filter":displacement_status,"support_type_filter":support_type,"projects":Project.objects.filter(active=True),"classes":[(f"P{i}",f"P{i}") for i in range(1,8)]+[(f"S{i}",f"S{i}") for i in range(1,7)],"statuses":Learner.Status.choices,"participant_categories":Learner.ParticipantCategory.choices,"displacement_statuses":Learner.DisplacementStatus.choices,"support_types":LearnerSupport._meta.get_field("support_type").choices,"kpis":kpis,"school_project_support":school_project_support})


@login_required
def supported_learners_by_school_project(request, school_id, project_id):
    school = get_object_or_404(School, pk=school_id)
    project = get_object_or_404(Project, pk=project_id)
    learners = (Learner.objects.filter(supports__project=project, enrolments__school=school, enrolments__is_current=True)
        .annotate(support_count=Count("supports", filter=Q(supports__project=project), distinct=True), last_support_date=Max("supports__distribution_date", filter=Q(supports__project=project)))
        .select_related("location").distinct().order_by("last_name", "first_name"))
    return render(request, "learners/supported_learners.html", {"school": school, "project": project, "learners": learners})


@login_required
def supported_learners_by_school(request, school_id):
    school = get_object_or_404(School, pk=school_id)
    supports, filters = _school_support_records(request, school)
    rows = _distribution_rows(supports, school)
    project_breakdown = list(LearnerSupport.objects.filter(learner__enrolments__school=school, learner__enrolments__is_current=True).values("project_id", "project__code", "project__name").annotate(total=Count("learner_id", distinct=True)).order_by("project__name"))
    return render(request, "learners/supported_learners_by_school.html", {
        "school": school,
        "project_breakdown": project_breakdown,
        "rows": rows,
        "projects": Project.objects.filter(active=True).order_by("name"),
        "support_types": LearnerSupport._meta.get_field("support_type").choices,
        "materials": LearnerSupport.SCHOLASTIC_MATERIAL_CHOICES,
        "filters": filters,
    })


def _school_support_records(request, school):
    filters = {key: request.GET.get(key, "").strip() for key in ("q", "project", "support_type", "scholastic_material", "class_level", "date_from", "date_to")}
    qs = LearnerSupport.objects.filter(learner__enrolments__school=school, learner__enrolments__is_current=True).select_related("learner", "project").distinct()
    if filters["q"]:
        qs = qs.filter(Q(learner__reg_no__icontains=filters["q"]) | Q(learner__first_name__icontains=filters["q"]) | Q(learner__last_name__icontains=filters["q"]))
    if filters["project"].isdigit(): qs = qs.filter(project_id=filters["project"])
    if filters["support_type"] in dict(LearnerSupport._meta.get_field("support_type").choices): qs = qs.filter(support_type=filters["support_type"])
    if filters["scholastic_material"] in dict(LearnerSupport.SCHOLASTIC_MATERIAL_CHOICES): qs = qs.filter(scholastic_material=filters["scholastic_material"])
    if filters["class_level"]: qs = qs.filter(learner__enrolments__class_level=filters["class_level"], learner__enrolments__is_current=True)
    if parse_date(filters["date_from"]): qs = qs.filter(distribution_date__gte=parse_date(filters["date_from"]))
    if parse_date(filters["date_to"]): qs = qs.filter(distribution_date__lte=parse_date(filters["date_to"]))
    return qs.order_by("learner__last_name", "learner__first_name", "-distribution_date"), filters


def _distribution_rows(supports, school):
    grouped = {}
    for support in supports:
        current_enrolment = support.learner.enrolments.filter(school=school, is_current=True).first()
        if support.learner_id not in grouped:
            grouped[support.learner_id] = {"learner": support.learner, "class_level": current_enrolment.class_level if current_enrolment else support.learner.current_class, "materials": {}}
        material = support.scholastic_material if support.support_type == "scholastic_materials" else ("sanitary_pads" if support.support_type == "sanitary_pads_mhm" else "")
        if material:
            grouped[support.learner_id]["materials"][material] = grouped[support.learner_id]["materials"].get(material, Decimal("0")) + (support.quantity or Decimal("0"))
    return list(grouped.values())


def _styled_distribution_document(Document, school, project, distribution_date, rows):
    """Build the EPICK paper distribution form without relying on an uploaded template."""
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.45)
    section.right_margin = Inches(0.45)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)

    logo_path = Path(settings.BASE_DIR) / "static" / "images" / "logo.png"
    if logo_path.exists():
        logo = document.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.add_run().add_picture(str(logo_path), width=Inches(0.62))
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("EPICK DISTRIBUTION FORM")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subrun = subtitle.add_run("Scholastic materials distribution and beneficiary acknowledgement")
    subrun.italic = True
    subrun.font.size = Pt(9)
    subrun.font.color.rgb = RGBColor(89, 89, 89)

    meta = document.add_table(rows=3, cols=2)
    meta.autofit = False
    meta_widths = [Inches(5.05), Inches(5.05)]
    metadata = [
        ("Activity", "Scholastic materials distribution", "Programme", project.code if project else "EPICK II"),
        ("School", school.name, "District", school.location.district if school.location else "________________"),
        ("Date", distribution_date or "________________", "Prepared by", "____________________________"),
    ]
    for row, values in zip(meta.rows, metadata):
        for cell, value, width in zip(row.cells, values, [meta_widths[0], meta_widths[0], meta_widths[1], meta_widths[1]]):
            cell.width = width
        row.cells[0].text = f"{values[0]}: {values[1]}"
        row.cells[1].text = f"{values[2]}: {values[3]}"
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(9)
            tc_pr = cell._tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "F2F6FA")
            tc_pr.append(shd)

    document.add_paragraph().paragraph_format.space_after = Pt(2)
    table = document.add_table(rows=1, cols=9)
    table.autofit = False
    table.style = "Table Grid"
    widths = [0.42, 3.05, 0.62, 0.68, 0.68, 0.68, 0.78, 0.68, 1.35]
    headers = ["S/N", "Name", "Class", "Soap", "Pads", "Pens", "Pencils", "Books", "Signature"]

    def shade(cell, fill, color="FFFFFF", bold=True):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string(color)
                run.bold = bold

    def set_cell(cell, text, centered=True):
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8)

    header = table.rows[0]
    for cell, header_text, width in zip(header.cells, headers, widths):
        cell.width = Inches(width)
        set_cell(cell, header_text)
        shade(cell, "1F4E79")
    header._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))

    for index, row_data in enumerate(rows, start=1):
        cells = table.add_row().cells
        materials = row_data["materials"]
        values = [index, row_data["learner"].full_name, row_data["class_level"] or "", materials.get("soap", ""), materials.get("sanitary_pads", ""), materials.get("pens", ""), materials.get("pencils", ""), materials.get("books", ""), ""]
        for cell, value, width in zip(cells, values, widths):
            cell.width = Inches(width)
            set_cell(cell, value, centered=(cell is not cells[1]))
        if index % 2 == 0:
            for cell in cells: shade(cell, "F2F6FA", color="183247", bold=False)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(8)
    note.add_run("Distribution officer: ").bold = True
    note.add_run("Please confirm quantities before collection and obtain each learner's signature.")
    return document


@login_required
def supported_learners_distribution_docx(request, school_id):
    school = get_object_or_404(School, pk=school_id)
    supports, filters = _school_support_records(request, school)
    rows = _distribution_rows(supports, school)
    try:
        from docx import Document
    except ImportError:
        return HttpResponse("The Word export dependency is not installed.", status=503)
    project = Project.objects.filter(pk=filters["project"]).first() if filters["project"].isdigit() else None
    document = _styled_distribution_document(Document, school, project, filters["date_from"], rows)
    output = BytesIO(); document.save(output); output.seek(0)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    response["Content-Disposition"] = f'attachment; filename="EPICK_distribution_{school.name.replace(" ", "_")}.docx"'
    return response


@login_required
@record_editor_required
def learner_create(request):
    form=LearnerForm(request.POST or None)
    if request.method=="POST" and form.is_valid():
        learner=form.save(commit=False); learner.created_by=request.user; learner.save()
        project_id=request.POST.get("project")
        if project_id:
            from projects.models import Project
            project=get_object_or_404(Project,pk=project_id)
            LearnerProject.objects.get_or_create(learner=learner,project=project,defaults={"start_date":learner.admission_date})
        messages.success(request,f"Learner registered successfully. KAPDA ID: {learner.reg_no}")
        return redirect("learners:detail",pk=learner.pk)
    from projects.models import Project
    return render(request,"learners/learner_registration_form.html",{"form":form,"projects":Project.objects.filter(active=True)})

@login_required
def learner_detail(request,pk):
    learner=get_object_or_404(Learner.objects.select_related("location"),pk=pk)
    return render(request,"learners/detail.html",{"learner":learner,"employment":getattr(learner, "employment", None),"enrolments":learner.enrolments.select_related("school","project","academic_year").order_by("-academic_year__start_date"),"attendance":learner.attendance_records.select_related("term").order_by("-term__start_date")[:8],"assessments":learner.assessments.select_related("term").order_by("-term__start_date")[:12],"supports":learner.supports.select_related("project","term").order_by("-distribution_date"),"uneb":learner.uneb_results.select_related("school","academic_year").order_by("-examination_year")})

@login_required
def status_history_list(request, learner_slug):
    learner = get_object_or_404(Learner, reg_no__iexact=learner_slug)
    history = learner.status_history.select_related("recorded_by").all()
    return render(request, "learners/status_history_list.html", {"learner": learner, "history": history})

@login_required
@record_editor_required
def status_history_create(request, learner_slug):
    learner = get_object_or_404(Learner, reg_no__iexact=learner_slug)
    form = LearnerStatusHistoryForm(request.POST or None, learner=learner)
    if request.method == "POST" and form.is_valid():
        history = form.save(commit=False)
        history.learner = learner
        history.recorded_by = request.user
        history.save()
        messages.success(request, f"{learner.full_name}'s status was updated to {history.get_new_status_display()}.")
        return redirect("learners:status_history_list", learner_slug=learner.reg_no.lower())
    return render(request, "learners/status_history_form.html", {"form": form, "learner": learner})

@login_required
def employment_list(request):
    q = request.GET.get("q", "").strip()
    records = LearnerEmployment.objects.select_related("learner").all()
    if q:
        records = records.filter(Q(learner__first_name__icontains=q) | Q(learner__last_name__icontains=q) | Q(employer__icontains=q) | Q(job_title__icontains=q))
    return render(request, "learners/employment_list.html", {"records": Paginator(records, 25).get_page(request.GET.get("page")), "query": q})

@login_required
def learner_employment_list(request, learner_slug):
    learner = get_object_or_404(Learner, reg_no__iexact=learner_slug)
    records = LearnerEmployment.objects.filter(learner=learner).select_related("learner")
    return render(request, "learners/employment_list.html", {
        "records": records,
        "query": "",
        "learner": learner,
    })

@login_required
@record_editor_required
def employment_create(request, pk):
    learner = get_object_or_404(Learner, pk=pk)
    form = LearnerEmploymentForm(request.POST or None, learner=learner, instance=getattr(learner, "employment", None))
    if request.method == "POST" and form.is_valid():
        employment = form.save(commit=False)
        employment.learner = learner
        employment.save()
        messages.success(request, "Employment outcome recorded successfully.")
        return redirect("learners:detail", pk=learner.pk)
    return render(request, "learners/employment_form.html", {"form": form, "learner": learner})

@login_required
@record_editor_required
def support_create(request, pk):
    learner = get_object_or_404(Learner, pk=pk)
    form = LearnerSupportForm(request.POST or None, learner=learner)
    if request.method == "POST" and form.is_valid():
        support = form.save(commit=False)
        support.learner = learner
        support.recorded_by = request.user
        support.save()
        messages.success(request, "Learner support recorded successfully.")
        return redirect("learners:detail", pk=learner.pk)
    return render(request, "learners/support_form.html", {"form": form, "learner": learner})

@login_required
@record_editor_required
def enrolment_create(request,pk):
    learner=get_object_or_404(Learner,pk=pk)
    latest_enrolment = learner.enrolments.select_related("school", "project").filter(is_current=True).first()
    current_year = AcademicYear.objects.filter(is_current=True).first()
    active_project = learner.project_links.filter(active=True).order_by("-start_date").first()
    enrolment_initial = {
        "school": None,
        "project": latest_enrolment.project_id if latest_enrolment else getattr(active_project, "project_id", None),
        "academic_year": current_year.pk if current_year else None,
        "class_level": learner.current_class or learner.class_of_admission,
    }
    form=EnrolmentForm(request.POST or None, initial=enrolment_initial)
    if request.method=="POST" and form.is_valid():
        school = form.cleaned_data["school"]
        academic_year = form.cleaned_data["academic_year"]
        if Enrolment.objects.filter(learner=learner, school=school).exists():
            messages.error(request, f"{learner.full_name} already has an enrolment history at {school}. A learner cannot be enrolled or transferred to the same school twice.")
            return redirect("learners:detail", pk=learner.pk)
        e=form.save(commit=False); e.learner=learner
        try:
            with transaction.atomic():
                list(Enrolment.objects.select_for_update().filter(learner=learner, is_current=True))
                e.is_current = True
                e.outcome = "enrolled"
                e.save()
        except IntegrityError:
            messages.error(request, f"The enrolment could not be saved because this learner already has that school or another current enrolment.")
            return redirect("learners:detail", pk=learner.pk)
        learner.current_class=e.class_level; learner.save(update_fields=["current_class","updated_at"])
        LearnerProject.objects.get_or_create(learner=learner,project=e.project,defaults={"start_date":e.academic_year.start_date})
        messages.success(request, "Transfer completed." if learner.enrolments.filter(outcome="transferred").exists() else "Enrolment saved.")
        return redirect("learners:detail",pk=learner.pk)
    return render(request,"learners/enrolment_create_form.html",{"form":form,"learner":learner})

@login_required
@record_editor_required
def retention(request):
    years=AcademicYear.objects.order_by("-start_date")
    if request.method=="POST":
        action=request.POST.get("action")
        if action=="retain":
            source_id=request.POST.get("source_year"); dest_id=request.POST.get("destination_year")
            class_level=request.POST.get("class_level"); school_id=request.POST.get("destination_school_id")
            source_school_id=request.POST.get("source_school_id")
            learner_ids=request.POST.getlist("learner_ids")
            if not (source_id and dest_id and class_level and school_id and learner_ids):
                messages.error(request,"Select source year, destination year, class, school and at least one learner.")
                return redirect("learners:retention")
            source=get_object_or_404(AcademicYear,pk=source_id); dest=get_object_or_404(AcademicYear,pk=dest_id)
            from core.models import School
            school=get_object_or_404(School,pk=school_id)
            if dest.start_date <= source.start_date:
                messages.error(request,"Destination academic year must be later than the source year.")
                return redirect("learners:retention")
            duplicate_enrolments = Enrolment.objects.filter(
                learner_id__in=learner_ids,
                school=school,
                academic_year=dest,
            ).select_related("learner")
            if duplicate_enrolments.exists():
                names = ", ".join(
                    enrolment.learner.full_name
                    for enrolment in duplicate_enrolments[:3]
                )
                suffix = "" if duplicate_enrolments.count() <= 3 else " and other selected learners"
                messages.error(
                    request,
                    f"Retention could not be completed. {names}{suffix} already has an enrolment "
                    f"at {school} for {dest}. No records were changed.",
                )
                return redirect("learners:retention")
            reused_school = []
            for current in Enrolment.objects.filter(pk__in=learner_ids, is_current=True):
                if current.school_id != school.pk and Enrolment.objects.filter(learner_id=current.learner_id, school_id=school.pk).exists():
                    reused_school.append(current.learner_id)
            if reused_school:
                messages.error(request, "One or more selected learners already have an enrolment history at the destination school. A learner cannot transfer to the same school twice.")
                return redirect("learners:retention")
            moved=0
            with transaction.atomic():
                qs=Enrolment.objects.select_for_update().filter(pk__in=learner_ids,is_current=True).select_related("learner","project","academic_year")
                if not qs.exists():
                    messages.error(request, "No selected current enrolments matched the retention request. Reload the class list, select learners, and try again.")
                    return redirect("learners:retention")
                for old in qs:
                    learner=old.learner
                    class_level = old.class_level
                    next_class=class_level
                    movement="repeat"; outcome="repeated"
                    if class_level.startswith("P") and int(class_level[1:])<7:
                        next_class=f"P{int(class_level[1:])+1}"; movement="retention"; outcome="retained"
                    elif class_level=="P7":
                        next_class="S1"; movement="transition"; outcome="completed"
                    elif class_level.startswith("S") and int(class_level[1:])<6:
                        next_class=f"S{int(class_level[1:])+1}"; movement="retention"; outcome="retained"
                    elif class_level=="S6":
                        next_class="S6"; movement="completion"; outcome="completed"
                    same_school = old.school_id == school.pk
                    new=None
                    if same_school:
                        # A learner has one enrolment history per school. Keep
                        # the same record current when progressing within the
                        # same school instead of creating a duplicate.
                        old.academic_year = dest
                        old.class_level = next_class
                        old.outcome = outcome
                        old.is_current = class_level != "S6"
                        old.save(update_fields=["academic_year", "class_level", "outcome", "is_current", "updated_at"])
                        new = old if old.is_current else None
                    else:
                        old.is_current=False; old.outcome="transferred"; old.save(update_fields=["is_current","outcome","updated_at"])
                        if Enrolment.objects.filter(learner=learner, school=school).exists():
                            raise IntegrityError(f"{learner.full_name} has already been enrolled at {school}.")
                        if class_level != "S6":
                            new=Enrolment.objects.create(learner=learner,school=school,project=old.project,academic_year=dest,class_level=next_class,is_current=True)
                    ClassMovement.objects.create(learner=learner,from_enrolment=old,to_enrolment=new,movement_type=movement,effective_date=dest.start_date,approved_by=request.user)
                    learner.current_class=next_class
                    if class_level=="S6": learner.status=Learner.Status.COMPLETED
                    else: learner.status=Learner.Status.ACTIVE
                    learner.save(update_fields=["current_class","status","updated_at"]); moved+=1
                AuditEvent.objects.create(actor=request.user,action="bulk_class_movement",model_label="learners.Enrolment",object_id="",summary=f"Moved {moved} learners from {class_level} in {source} to next class in {dest}.")
            messages.success(request,f"Successfully processed {moved} learners. Review the resulting enrolments before closing the academic year.")
            return redirect("learners:retention")
    from core.models import School
    source_id=request.GET.get("source_year") or (str(years[0].pk) if years else "")
    class_level=request.GET.get("class_level","P1"); school_id=request.GET.get("school","")
    enrolments=Enrolment.objects.filter(is_current=True,class_level=class_level).select_related("learner","school","project","academic_year")
    if source_id: enrolments=enrolments.filter(academic_year_id=source_id)
    if school_id: enrolments=enrolments.filter(school_id=school_id)
    return render(request,"learners/retention.html",{"years":years,"schools":School.objects.filter(active=True),"source_id":source_id,"class_level":class_level,"school_id":school_id,"enrolments":enrolments,"classes":[(f"P{i}",f"P{i}") for i in range(1,8)]+[(f"S{i}",f"S{i}") for i in range(1,7)],"statuses":Learner.Status.choices})
